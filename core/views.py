from django.http import JsonResponse
from datetime import datetime

from django.contrib import messages
from django.db.models import F
from django.db import models
from django.http import HttpResponse
from django.shortcuts import render, redirect, get_object_or_404

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph
)

from .models import (
    Cliente,
    Proveedor,
    Producto,
    Empleado,
)

from .forms import EmpleadoForm


# ==========================================
# LOGIN
# ==========================================

def login(request):

    return render(
        request,
        "login/index.html"
    )


# ==========================================
# DASHBOARD
# ==========================================

def dashboard(request):

    context = {

        "total_clientes": Cliente.objects.count(),
        "total_proveedores": Proveedor.objects.count(),
        "total_productos": Producto.objects.count(),
        "total_empleados": Empleado.objects.count(),

        "form": EmpleadoForm(),
        "proveedores": Proveedor.objects.all(),

        "clientes": Cliente.objects.all(),
        "productos": Producto.objects.all(),

    }

    return render(request, "dashboard/dashboard.html", context)

# ==========================================
# EMPLEADOS
# ==========================================

def empleados(request):

    if request.method == "POST":

        form = EmpleadoForm(request.POST)

        if form.is_valid():

            form.save()

            # Si la petición viene por AJAX (fetch desde el Dashboard)
            if request.headers.get("X-Requested-With") == "XMLHttpRequest":
                return JsonResponse({
                    "success": True,
                    "message": "Empleado agregado correctamente."
                })

            messages.success(
                request,
                "Empleado agregado correctamente."
            )

            return redirect("empleados")

        else:
            # Formulario inválido y viene por AJAX
            if request.headers.get("X-Requested-With") == "XMLHttpRequest":
                return JsonResponse({
                    "success": False,
                    "errors": form.errors
                }, status=400)

    else:

        form = EmpleadoForm()

    lista_empleados = Empleado.objects.all()

    context = {

        "form": form,

        "empleados": lista_empleados,

        "total_empleados":
            lista_empleados.count(),

        "empleados_activos":
            lista_empleados.filter(
                estado="Activo"
            ).count(),

        "empleados_vacaciones":
            lista_empleados.filter(
                estado="Vacaciones"
            ).count(),

        "empleados_nuevos":
            lista_empleados.count(),

    }

    return render(
        request,
        "empleados/empleados.html",
        context
    )
# ==========================================
# EDITAR EMPLEADO
# ==========================================

def editar_empleado(request, id):

    empleado = get_object_or_404(
        Empleado,
        id=id
    )

    if request.method == "POST":

        form = EmpleadoForm(
            request.POST,
            instance=empleado
        )

        if form.is_valid():

            form.save()

            messages.success(
                request,
                "Empleado actualizado correctamente."
            )

            return redirect("empleados")

    else:

        form = EmpleadoForm(
            instance=empleado
        )

    return render(
        request,
        "empleados/editar_empleado.html",
        {
            "form": form,
            "empleado": empleado
        }
    )


# ==========================================
# ELIMINAR EMPLEADO
# ==========================================

def eliminar_empleado(request, id):

    empleado = get_object_or_404(
        Empleado,
        id=id
    )

    empleado.delete()

    messages.success(
        request,
        "Empleado eliminado correctamente."
    )

    return redirect("empleados")
# ==========================================
# PRODUCTOS
# ==========================================

def productos(request):

    lista_productos = Producto.objects.all()

    # Productos cuyo stock está por debajo o igual al mínimo
    productos_stock_bajo = lista_productos.filter(
        stock__lte=F("stock_minimo")
    )

    context = {
        "productos": lista_productos,
        "total_productos": lista_productos.count(),
        "productos_stock_bajo": productos_stock_bajo.count(),
    }

    return render(
        request,
        "productos/productos.html",
        context
    )
# ==========================================
# EXPORTAR EMPLEADOS EXCEL
# ==========================================

def exportar_empleados_excel(request):

    wb = Workbook()

    ws = wb.active

    ws.title = "Empleados"

    encabezados = [
        "Documento",
        "Nombres",
        "Apellidos",
        "Correo",
        "Teléfono",
        "Cargo",
        "Área",
        "Estado",
        "Fecha de ingreso",
    ]

    color_fondo = PatternFill(
        start_color="1F4E78",
        end_color="1F4E78",
        fill_type="solid"
    )

    fuente = Font(
        color="FFFFFF",
        bold=True
    )

    alineacion = Alignment(
        horizontal="center",
        vertical="center"
    )

    for columna, titulo in enumerate(
        encabezados,
        start=1
    ):

        celda = ws.cell(
            row=1,
            column=columna
        )

        celda.value = titulo
        celda.fill = color_fondo
        celda.font = fuente
        celda.alignment = alineacion

    empleados = Empleado.objects.all()

    fila = 2

    for empleado in empleados:

        ws.cell(
            row=fila,
            column=1
        ).value = empleado.documento

        ws.cell(
            row=fila,
            column=2
        ).value = empleado.nombres

        ws.cell(
            row=fila,
            column=3
        ).value = empleado.apellidos

        ws.cell(
            row=fila,
            column=4
        ).value = empleado.correo

        ws.cell(
            row=fila,
            column=5
        ).value = empleado.telefono

        ws.cell(
            row=fila,
            column=6
        ).value = empleado.cargo

        ws.cell(
            row=fila,
            column=7
        ).value = empleado.area

        ws.cell(
            row=fila,
            column=8
        ).value = empleado.estado

        ws.cell(
            row=fila,
            column=9
        ).value = empleado.fecha_ingreso

        fila += 1

    fecha = datetime.now().strftime(
        "%d-%m-%Y"
    )

    response = HttpResponse(
        content_type=
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    response["Content-Disposition"] = (
        f'attachment; filename="Empleados_{fecha}.xlsx"'
    )

    wb.save(response)

    return response


# ==========================================
# EXPORTAR EMPLEADOS PDF
# ==========================================

def exportar_empleados_pdf(request):

    response = HttpResponse(
        content_type="application/pdf"
    )

    fecha_actual = datetime.now().strftime(
        "%d-%m-%Y"
    )

    response["Content-Disposition"] = (
        f'attachment; filename="Empleados_{fecha_actual}.pdf"'
    )

    # ======================================================
    # DOCUMENTO
    # ======================================================

    documento_pdf = SimpleDocTemplate(
        response,
        pagesize=letter,
        rightMargin=25,
        leftMargin=25,
        topMargin=30,
        bottomMargin=30
    )


    # ======================================================
    # ESTILOS
    # ======================================================

    estilos = getSampleStyleSheet()


    titulo = ParagraphStyle(
        "TituloEmpleados",
        parent=estilos["Title"],
        fontName="Helvetica-Bold",
        fontSize=18,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#1F4E78"),
        spaceAfter=6
    )


    subtitulo = ParagraphStyle(
        "SubtituloEmpleados",
        parent=estilos["Normal"],
        fontName="Helvetica",
        fontSize=9,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#666666"),
        spaceAfter=15
    )


    texto_tabla = ParagraphStyle(
        "TextoTabla",
        parent=estilos["Normal"],
        fontName="Helvetica",
        fontSize=7,
        leading=9,
        alignment=TA_CENTER
    )


    texto_tabla_izquierda = ParagraphStyle(
        "TextoTablaIzquierda",
        parent=estilos["Normal"],
        fontName="Helvetica",
        fontSize=7,
        leading=9,
        alignment=TA_LEFT
    )


    # ======================================================
    # ENCABEZADO
    # ======================================================

    elementos = []


    elementos.append(
        Paragraph(
            "MINING STAR ERP",
            titulo
        )
    )


    elementos.append(
        Paragraph(
            "Listado de Empleados",
            subtitulo
        )
    )


    elementos.append(
        Paragraph(
            f"Fecha de generación: {fecha_actual}",
            subtitulo
        )
    )


    # ======================================================
    # ENCABEZADOS DE LA TABLA
    # ======================================================

    datos = [[

        Paragraph(
            "<b>Documento</b>",
            texto_tabla
        ),

        Paragraph(
            "<b>Nombre</b>",
            texto_tabla
        ),

        Paragraph(
            "<b>Cargo</b>",
            texto_tabla
        ),

        Paragraph(
            "<b>Área</b>",
            texto_tabla
        ),

        Paragraph(
            "<b>Correo</b>",
            texto_tabla
        ),

        Paragraph(
            "<b>Teléfono</b>",
            texto_tabla
        ),

        Paragraph(
            "<b>Estado</b>",
            texto_tabla
        ),

    ]]


    # ======================================================
    # EMPLEADOS
    # ======================================================

    empleados = Empleado.objects.all()


    for empleado in empleados:

        nombre_completo = (
            f"{empleado.nombres} "
            f"{empleado.apellidos}"
        )


        datos.append([

            Paragraph(
                str(empleado.documento),
                texto_tabla
            ),

            Paragraph(
                nombre_completo,
                texto_tabla_izquierda
            ),

            Paragraph(
                str(empleado.cargo),
                texto_tabla_izquierda
            ),

            Paragraph(
                str(empleado.area),
                texto_tabla_izquierda
            ),

            Paragraph(
                str(empleado.correo),
                texto_tabla_izquierda
            ),

            Paragraph(
                str(empleado.telefono),
                texto_tabla
            ),

            Paragraph(
                str(empleado.estado),
                texto_tabla
            ),

        ])


    # ======================================================
    # TABLA
    # ======================================================

    tabla = Table(
        datos,
        repeatRows=1,
        colWidths=[
            70,
            100,
            80,
            70,
            120,
            65,
            65,
        ]
    )


    # ======================================================
    # ESTILO DE TABLA
    # ======================================================

    tabla.setStyle(
        TableStyle([

            # ENCABEZADO

            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.HexColor("#1F4E78")
            ),

            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.white
            ),

            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold"
            ),

            (
                "VALIGN",
                (0, 0),
                (-1, -1),
                "MIDDLE"
            ),

            (
                "ALIGN",
                (0, 0),
                (-1, -1),
                "CENTER"
            ),

            # BORDES

            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.5,
                colors.HexColor("#B7B7B7")
            ),

            # ESPACIADO

            (
                "TOPPADDING",
                (0, 0),
                (-1, -1),
                6
            ),

            (
                "BOTTOMPADDING",
                (0, 0),
                (-1, -1),
                6
            ),

            # FILAS

            (
                "ROWBACKGROUNDS",
                (0, 1),
                (-1, -1),
                [
                    colors.white,
                    colors.HexColor("#F3F6F9")
                ]
            ),

        ])
    )


    # ======================================================
    # CONSTRUIR PDF
    # ======================================================

    elementos.append(tabla)


    documento_pdf.build(
        elementos
    )


    return response

# ==========================================
# CLIENTES
# ==========================================

def clientes(request):

    if request.method == "POST":

        Cliente.objects.create(

            documento=request.POST.get("documento"),

            nombre=request.POST.get("nombre"),

            correo=request.POST.get("correo"),

            telefono=request.POST.get("telefono"),

            empresa=request.POST.get("empresa"),

        )

        # Si la petición viene por AJAX (fetch desde el Dashboard)
        if request.headers.get("X-Requested-With") == "XMLHttpRequest":
            return JsonResponse({
                "success": True,
                "message": "Cliente agregado correctamente."
            })

        messages.success(
            request,
            "Cliente agregado correctamente."
        )

        return redirect("clientes")

    lista_clientes = Cliente.objects.all()

    context = {

        "clientes": lista_clientes,

        "total_clientes":
            lista_clientes.count(),

    }

    return render(
        request,
        "clientes/clientes.html",
        context
    )


# ==========================================
# EDITAR CLIENTE
# ==========================================

def editar_cliente(request, id):

    cliente = get_object_or_404(
        Cliente,
        id=id
    )

    if request.method == "POST":

        cliente.documento = request.POST.get(
            "documento"
        )

        cliente.nombre = request.POST.get(
            "nombre"
        )

        cliente.correo = request.POST.get(
            "correo"
        )

        cliente.telefono = request.POST.get(
            "telefono"
        )

        cliente.empresa = request.POST.get(
            "empresa"
        )

        cliente.save()

        messages.success(
            request,
            "Cliente actualizado correctamente."
        )

        return redirect("clientes")

    return render(
        request,
        "clientes/editar_cliente.html",
        {
            "cliente": cliente
        }
    )


# ==========================================
# ELIMINAR CLIENTE
# ==========================================

def eliminar_cliente(request, id):

    cliente = get_object_or_404(
        Cliente,
        id=id
    )

    cliente.delete()

    messages.success(
        request,
        "Cliente eliminado correctamente."
    )

    return redirect("clientes")


# ==========================================
# EXPORTAR CLIENTES EXCEL
# ==========================================

def exportar_clientes_excel(request):

    wb = Workbook()

    ws = wb.active

    ws.title = "Clientes"

    encabezados = [
        "Documento",
        "Nombre",
        "Correo",
        "Teléfono",
        "Empresa"
    ]

    color_fondo = PatternFill(
        start_color="1F4E78",
        end_color="1F4E78",
        fill_type="solid"
    )

    fuente = Font(
        color="FFFFFF",
        bold=True
    )

    alineacion = Alignment(
        horizontal="center",
        vertical="center"
    )

    for columna, titulo in enumerate(
        encabezados,
        start=1
    ):

        celda = ws.cell(
            row=1,
            column=columna
        )

        celda.value = titulo

        celda.fill = color_fondo

        celda.font = fuente

        celda.alignment = alineacion

    clientes = Cliente.objects.all()

    fila = 2

    for cliente in clientes:

        ws.cell(
            row=fila,
            column=1
        ).value = cliente.documento

        ws.cell(
            row=fila,
            column=2
        ).value = cliente.nombre

        ws.cell(
            row=fila,
            column=3
        ).value = cliente.correo

        ws.cell(
            row=fila,
            column=4
        ).value = cliente.telefono

        ws.cell(
            row=fila,
            column=5
        ).value = cliente.empresa

        fila += 1

    fecha = datetime.now().strftime(
        "%d-%m-%Y"
    )

    response = HttpResponse(
        content_type=
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    response["Content-Disposition"] = (
        f'attachment; filename="Clientes_{fecha}.xlsx"'
    )

    wb.save(response)

    return response


# ==========================================
# EXPORTAR CLIENTES PDF
# ==========================================

def exportar_clientes_pdf(request):

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        SimpleDocTemplate,
        Table,
        TableStyle,
        Paragraph,
        Spacer
    )
    from datetime import datetime

    # ==========================================
    # RESPUESTA
    # ==========================================

    response = HttpResponse(
        content_type="application/pdf"
    )

    fecha_archivo = datetime.now().strftime(
        "%d-%m-%Y"
    )

    response["Content-Disposition"] = (
        f'attachment; filename="Reporte_Clientes_{fecha_archivo}.pdf"'
    )

    # ==========================================
    # DOCUMENTO
    # ==========================================

    documento = SimpleDocTemplate(
        response,
        pagesize=landscape(letter),
        rightMargin=12 * mm,
        leftMargin=12 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm
    )

    # ==========================================
    # ESTILOS
    # ==========================================

    estilos = getSampleStyleSheet()

    titulo = ParagraphStyle(
        "TituloClientes",
        parent=estilos["Title"],
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=24,
        alignment=TA_CENTER,
        spaceAfter=8
    )

    subtitulo = ParagraphStyle(
        "SubtituloClientes",
        parent=estilos["Normal"],
        fontName="Helvetica",
        fontSize=10,
        leading=14,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#64748B"),
        spaceAfter=15
    )

    texto_tabla = ParagraphStyle(
        "TextoTablaClientes",
        parent=estilos["Normal"],
        fontName="Helvetica",
        fontSize=8,
        leading=10,
        alignment=TA_LEFT
    )

    texto_centro = ParagraphStyle(
        "TextoCentroClientes",
        parent=texto_tabla,
        alignment=TA_CENTER
    )

    # ==========================================
    # CLIENTES
    # ==========================================

    clientes = Cliente.objects.all()

    total_clientes = clientes.count()

    # ==========================================
    # FECHA
    # ==========================================

    fecha = datetime.now().strftime(
        "%d/%m/%Y %H:%M"
    )

    # ==========================================
    # CONTENIDO
    # ==========================================

    elementos = []

    # ==========================================
    # TÍTULO
    # ==========================================

    elementos.append(
        Paragraph(
            "MINING STAR ERP",
            titulo
        )
    )

    elementos.append(
        Paragraph(
            "REPORTE DE CLIENTES",
            subtitulo
        )
    )

    # ==========================================
    # FECHA
    # ==========================================

    elementos.append(
        Paragraph(
            f"<b>Fecha de generación:</b> {fecha}",
            texto_centro
        )
    )

    elementos.append(
        Spacer(1, 8)
    )

    # ==========================================
    # TOTAL CLIENTES
    # ==========================================

    elementos.append(
        Paragraph(
            f"<b>Total de clientes:</b> {total_clientes}",
            texto_centro
        )
    )

    elementos.append(
        Spacer(1, 15)
    )

    # ==========================================
    # ENCABEZADOS
    # ==========================================

    datos = [

        [
            Paragraph(
                "<b>Documento</b>",
                texto_centro
            ),

            Paragraph(
                "<b>Nombre</b>",
                texto_centro
            ),

            Paragraph(
                "<b>Empresa</b>",
                texto_centro
            ),

            Paragraph(
                "<b>Correo</b>",
                texto_centro
            ),

            Paragraph(
                "<b>Teléfono</b>",
                texto_centro
            ),

            Paragraph(
                "<b>Estado</b>",
                texto_centro
            ),
        ]

    ]

    # ==========================================
    # FILAS
    # ==========================================

    for cliente in clientes:

        datos.append(

            [

                Paragraph(
                    str(cliente.documento),
                    texto_centro
                ),

                Paragraph(
                    str(cliente.nombre),
                    texto_tabla
                ),

                Paragraph(
                    str(cliente.empresa or "-"),
                    texto_tabla
                ),

                Paragraph(
                    str(cliente.correo),
                    texto_tabla
                ),

                Paragraph(
                    str(cliente.telefono),
                    texto_centro
                ),

                Paragraph(
                    "Activo",
                    texto_centro
                ),

            ]

        )

    # ==========================================
    # TABLA
    # ==========================================

    tabla = Table(
        datos,
        repeatRows=1,
        colWidths=[
            35 * mm,
            55 * mm,
            55 * mm,
            65 * mm,
            35 * mm,
            30 * mm,
        ]
    )

    # ==========================================
    # ESTILO TABLA
    # ==========================================

    estilo_tabla = [

        # ENCABEZADO

        (
            "BACKGROUND",
            (0, 0),
            (-1, 0),
            colors.HexColor("#1F4E78")
        ),

        (
            "TEXTCOLOR",
            (0, 0),
            (-1, 0),
            colors.white
        ),

        (
            "FONTNAME",
            (0, 0),
            (-1, 0),
            "Helvetica-Bold"
        ),

        # CUERPO

        (
            "FONTNAME",
            (0, 1),
            (-1, -1),
            "Helvetica"
        ),

        (
            "FONTSIZE",
            (0, 1),
            (-1, -1),
            8
        ),

        # ALINEACIÓN

        (
            "VALIGN",
            (0, 0),
            (-1, -1),
            "MIDDLE"
        ),

        # BORDES

        (
            "GRID",
            (0, 0),
            (-1, -1),
            0.5,
            colors.HexColor("#CBD5E1")
        ),

        # FILAS ALTERNADAS

        (
            "ROWBACKGROUNDS",
            (0, 1),
            (-1, -1),
            [
                colors.white,
                colors.HexColor("#F8FAFC")
            ]
        ),

        # ESPACIADO

        (
            "TOPPADDING",
            (0, 0),
            (-1, -1),
            7
        ),

        (
            "BOTTOMPADDING",
            (0, 0),
            (-1, -1),
            7
        ),

        (
            "LEFTPADDING",
            (0, 0),
            (-1, -1),
            5
        ),

        (
            "RIGHTPADDING",
            (0, 0),
            (-1, -1),
            5
        ),

    ]

    tabla.setStyle(
        TableStyle(estilo_tabla)
    )

    # ==========================================
    # AGREGAR TABLA
    # ==========================================

    elementos.append(
        tabla
    )

    elementos.append(
        Spacer(1, 15)
    )

    # ==========================================
    # PIE DEL REPORTE
    # ==========================================

    elementos.append(
        Paragraph(
            "Mining Star ERP | Sistema Integral de Gestión Minera",
            subtitulo
        )
    )

    # ==========================================
    # GENERAR PDF
    # ==========================================

    documento.build(
        elementos
    )

    return response

# ==========================================
# NUEVO PRODUCTO
# ==========================================

def nuevo_producto(request):

    proveedores = Proveedor.objects.all()

    if request.method == "POST":

        Producto.objects.create(
            codigo=request.POST.get("codigo"),
            nombre=request.POST.get("nombre"),
            categoria=request.POST.get("categoria"),
            proveedor_id=request.POST.get("proveedor"),
            precio_compra=request.POST.get("precio_compra"),
            precio_venta=request.POST.get("precio_venta"),
            stock=request.POST.get("stock"),
            stock_minimo=request.POST.get("stock_minimo"),
        )

        # Si la petición viene por AJAX (fetch desde el Dashboard)
        if request.headers.get("X-Requested-With") == "XMLHttpRequest":
            return JsonResponse({
                "success": True,
                "message": "Producto registrado correctamente."
            })

        messages.success(
            request,
            "Producto registrado correctamente."
        )

        return redirect("productos")

    return render(
        request,
        "productos/nuevo_producto.html",
        {
            "proveedores": proveedores
        }
    )


# ==========================================
# EDITAR PRODUCTO
# ==========================================

def editar_producto(request, id):

    producto = get_object_or_404(
        Producto,
        id=id
    )

    if request.method == "POST":

        producto.codigo = request.POST.get("codigo")
        producto.nombre = request.POST.get("nombre")
        producto.categoria = request.POST.get("categoria")
        producto.proveedor_id = request.POST.get("proveedor")

        producto.precio_compra = request.POST.get(
            "precio_compra"
        )

        producto.precio_venta = request.POST.get(
            "precio_venta"
        )

        producto.stock = request.POST.get(
            "stock"
        )

        producto.stock_minimo = request.POST.get(
            "stock_minimo"
        )

        producto.save()

        messages.success(
            request,
            "Producto actualizado correctamente."
        )

        return redirect("productos")

    proveedores = Proveedor.objects.all()

    context = {
        "producto": producto,
        "proveedores": proveedores,
    }

    return render(
        request,
        "productos/editar_producto.html",
        context
    )

# ==========================================
# ELIMINAR PRODUCTO
# ==========================================

def eliminar_producto(request, id):

    producto = get_object_or_404(
        Producto,
        id=id
    )

    producto.delete()

    messages.success(
        request,
        "Producto eliminado correctamente."
    )

    return redirect("productos")


# ==========================================
# EXPORTAR PRODUCTOS EXCEL
# ==========================================

def exportar_productos_excel(request):

    wb = Workbook()

    ws = wb.active

    ws.title = "Productos"

    encabezados = [
        "Código",
        "Producto",
        "Categoría",
        "Proveedor",
        "Precio Compra",
        "Precio Venta",
        "Stock",
        "Stock Mínimo",
    ]

    color_fondo = PatternFill(
        start_color="1F4E78",
        end_color="1F4E78",
        fill_type="solid"
    )

    fuente = Font(
        color="FFFFFF",
        bold=True
    )

    alineacion = Alignment(
        horizontal="center",
        vertical="center"
    )

    for columna, titulo in enumerate(
        encabezados,
        start=1
    ):

        celda = ws.cell(
            row=1,
            column=columna
        )

        celda.value = titulo
        celda.fill = color_fondo
        celda.font = fuente
        celda.alignment = alineacion

    productos = Producto.objects.select_related(
        "proveedor"
    ).all()

    fila = 2

    for producto in productos:

        datos = [
            producto.codigo,
            producto.nombre,
            producto.categoria,
            producto.proveedor.empresa,
            producto.precio_compra,
            producto.precio_venta,
            producto.stock,
            producto.stock_minimo,
        ]

        for columna, valor in enumerate(
            datos,
            start=1
        ):

            ws.cell(
                row=fila,
                column=columna
            ).value = valor

        fila += 1

    fecha = datetime.now().strftime(
        "%d-%m-%Y"
    )

    response = HttpResponse(
        content_type=
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    response["Content-Disposition"] = (
        f'attachment; filename="Productos_{fecha}.xlsx"'
    )

    wb.save(response)

    return response


# ==========================================
# EXPORTAR PRODUCTOS PDF
# ==========================================

def exportar_productos_pdf(request):

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        SimpleDocTemplate,
        Table,
        TableStyle,
        Paragraph,
        Spacer
    )
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from datetime import datetime

    # ==========================================
    # RESPUESTA
    # ==========================================

    response = HttpResponse(
        content_type="application/pdf"
    )

    response["Content-Disposition"] = (
        'attachment; filename="Reporte_Productos.pdf"'
    )

    # ==========================================
    # DOCUMENTO
    # ==========================================

    documento = SimpleDocTemplate(
        response,
        pagesize=landscape(letter),
        rightMargin=12 * mm,
        leftMargin=12 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm
    )

    # ==========================================
    # ESTILOS
    # ==========================================

    estilos = getSampleStyleSheet()

    titulo = ParagraphStyle(
        "Titulo",
        parent=estilos["Title"],
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=24,
        alignment=TA_CENTER,
        spaceAfter=8
    )

    subtitulo = ParagraphStyle(
        "Subtitulo",
        parent=estilos["Normal"],
        fontName="Helvetica",
        fontSize=10,
        leading=14,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#64748B"),
        spaceAfter=15
    )

    texto_tabla = ParagraphStyle(
        "TextoTabla",
        parent=estilos["Normal"],
        fontName="Helvetica",
        fontSize=8,
        leading=10,
        alignment=TA_LEFT
    )

    texto_centro = ParagraphStyle(
        "TextoCentro",
        parent=texto_tabla,
        alignment=TA_CENTER
    )

    # ==========================================
    # PRODUCTOS
    # ==========================================

    productos = Producto.objects.select_related(
        "proveedor"
    ).all()

    total_productos = productos.count()

    stock_bajo = productos.filter(
        stock__lte=models.F("stock_minimo")
    ).count()

    # ==========================================
    # FECHA
    # ==========================================

    fecha = datetime.now().strftime(
        "%d/%m/%Y %H:%M"
    )

    # ==========================================
    # CONTENIDO
    # ==========================================

    elementos = []

    elementos.append(
        Paragraph(
            "MINING STAR ERP",
            titulo
        )
    )

    elementos.append(
        Paragraph(
            "REPORTE DE PRODUCTOS",
            subtitulo
        )
    )

    elementos.append(
        Paragraph(
            f"<b>Fecha de generación:</b> {fecha}",
            texto_centro
        )
    )

    elementos.append(
        Spacer(1, 8)
    )

    elementos.append(
        Paragraph(
            f"<b>Total de productos:</b> {total_productos} "
            f"&nbsp;&nbsp;&nbsp;&nbsp;"
            f"<b>Stock bajo:</b> {stock_bajo}",
            texto_centro
        )
    )

    elementos.append(
        Spacer(1, 15)
    )

    # ==========================================
    # ENCABEZADOS
    # ==========================================

    datos = [

        [
            Paragraph("<b>Código</b>", texto_centro),
            Paragraph("<b>Producto</b>", texto_centro),
            Paragraph("<b>Categoría</b>", texto_centro),
            Paragraph("<b>Proveedor</b>", texto_centro),
            Paragraph("<b>Precio Compra</b>", texto_centro),
            Paragraph("<b>Precio Venta</b>", texto_centro),
            Paragraph("<b>Stock</b>", texto_centro),
            Paragraph("<b>Stock Mínimo</b>", texto_centro),
        ]

    ]

    # ==========================================
    # FILAS
    # ==========================================

    for producto in productos:

        es_stock_bajo = (
            producto.stock <= producto.stock_minimo
        )

        stock_texto = str(producto.stock)

        if es_stock_bajo:

            stock_texto = (
                f'<font color="#DC2626">'
                f'<b>⚠ {producto.stock}</b>'
                f'</font>'
            )

        datos.append(

            [

                Paragraph(
                    str(producto.codigo),
                    texto_centro
                ),

                Paragraph(
                    str(producto.nombre),
                    texto_tabla
                ),

                Paragraph(
                    str(producto.categoria),
                    texto_tabla
                ),

                Paragraph(
                    str(producto.proveedor),
                    texto_tabla
                ),

                Paragraph(
                    f"${producto.precio_compra:,.2f}",
                    texto_centro
                ),

                Paragraph(
                    f"${producto.precio_venta:,.2f}",
                    texto_centro
                ),

                Paragraph(
                    stock_texto,
                    texto_centro
                ),

                Paragraph(
                    str(producto.stock_minimo),
                    texto_centro
                ),

            ]

        )

    # ==========================================
    # TABLA
    # ==========================================

    tabla = Table(
        datos,
        repeatRows=1,
        colWidths=[
            25 * mm,
            43 * mm,
            30 * mm,
            38 * mm,
            30 * mm,
            30 * mm,
            22 * mm,
            28 * mm,
        ]
    )

    # ==========================================
    # ESTILO TABLA
    # ==========================================

    estilo_tabla = [

        (
            "BACKGROUND",
            (0, 0),
            (-1, 0),
            colors.HexColor("#1F4E78")
        ),

        (
            "TEXTCOLOR",
            (0, 0),
            (-1, 0),
            colors.white
        ),

        (
            "FONTNAME",
            (0, 0),
            (-1, 0),
            "Helvetica-Bold"
        ),

        (
            "VALIGN",
            (0, 0),
            (-1, -1),
            "MIDDLE"
        ),

        (
            "GRID",
            (0, 0),
            (-1, -1),
            0.5,
            colors.HexColor("#CBD5E1")
        ),

        (
            "ROWBACKGROUNDS",
            (0, 1),
            (-1, -1),
            [
                colors.white,
                colors.HexColor("#F8FAFC")
            ]
        ),

        (
            "TOPPADDING",
            (0, 0),
            (-1, -1),
            7
        ),

        (
            "BOTTOMPADDING",
            (0, 0),
            (-1, -1),
            7
        ),

        (
            "LEFTPADDING",
            (0, 0),
            (-1, -1),
            5
        ),

        (
            "RIGHTPADDING",
            (0, 0),
            (-1, -1),
            5
        ),

    ]

    # ==========================================
    # RESALTAR STOCK BAJO
    # ==========================================

    for indice, producto in enumerate(
        productos,
        start=1
    ):

        if producto.stock <= producto.stock_minimo:

            estilo_tabla.append(
                (
                    "BACKGROUND",
                    (6, indice),
                    (6, indice),
                    colors.HexColor("#FEE2E2")
                )
            )

    tabla.setStyle(
        TableStyle(estilo_tabla)
    )

    elementos.append(tabla)

    elementos.append(
        Spacer(1, 15)
    )

    # ==========================================
    # PIE DEL REPORTE
    # ==========================================

    elementos.append(
        Paragraph(
            "Mining Star ERP | Sistema Integral de Gestión Minera",
            subtitulo
        )
    )

    # ==========================================
    # GENERAR PDF
    # ==========================================

    documento.build(elementos)

    return response
# ==========================================
# PROVEEDORES
# ==========================================

def proveedores(request):

    if request.method == "POST":

        Proveedor.objects.create(
            nit=request.POST.get("nit"),
            empresa=request.POST.get("empresa"),
            contacto=request.POST.get("contacto"),
            correo=request.POST.get("correo"),
            telefono=request.POST.get("telefono"),
            ciudad=request.POST.get("ciudad"),
            direccion=request.POST.get("direccion"),
        )

        # Si la petición viene por AJAX (fetch desde el Dashboard)
        if request.headers.get("X-Requested-With") == "XMLHttpRequest":
            return JsonResponse({
                "success": True,
                "message": "Proveedor agregado correctamente."
            })

        messages.success(
            request,
            "Proveedor agregado correctamente."
        )

        return redirect("proveedores")

    lista_proveedores = Proveedor.objects.all()

    context = {
        "proveedores": lista_proveedores,
        "total_proveedores": lista_proveedores.count(),
    }

    return render(
        request,
        "proveedores/proveedores.html",
        context
    )


# ==========================================
# EDITAR PROVEEDOR
# ==========================================

def editar_proveedor(request, id):

    proveedor = get_object_or_404(
        Proveedor,
        id=id
    )

    if request.method == "POST":

        proveedor.nit = request.POST.get("nit")
        proveedor.empresa = request.POST.get("empresa")
        proveedor.contacto = request.POST.get("contacto")
        proveedor.correo = request.POST.get("correo")
        proveedor.telefono = request.POST.get("telefono")
        proveedor.ciudad = request.POST.get("ciudad")
        proveedor.direccion = request.POST.get("direccion")

        proveedor.save()

        messages.success(
            request,
            "Proveedor actualizado correctamente."
        )

        return redirect("proveedores")

    return render(
        request,
        "proveedores/editar_proveedor.html",
        {
            "proveedor": proveedor
        }
    )


# ==========================================
# ELIMINAR PROVEEDOR
# ==========================================

def eliminar_proveedor(request, id):

    proveedor = get_object_or_404(
        Proveedor,
        id=id
    )

    proveedor.delete()

    messages.success(
        request,
        "Proveedor eliminado correctamente."
    )

    return redirect("proveedores")


# ==========================================
# EXPORTAR PROVEEDORES EXCEL
# ==========================================

def exportar_proveedores_excel(request):

    wb = Workbook()

    ws = wb.active
    ws.title = "Proveedores"

    encabezados = [
        "NIT",
        "Empresa",
        "Contacto",
        "Correo",
        "Teléfono",
        "Ciudad",
        "Dirección"
    ]

    for columna, titulo in enumerate(
        encabezados,
        start=1
    ):

        celda = ws.cell(
            row=1,
            column=columna
        )

        celda.value = titulo
        celda.font = Font(
            color="FFFFFF",
            bold=True
        )

        celda.fill = PatternFill(
            start_color="1F4E78",
            end_color="1F4E78",
            fill_type="solid"
        )

        celda.alignment = Alignment(
            horizontal="center"
        )

    proveedores = Proveedor.objects.all()

    fila = 2

    for proveedor in proveedores:

        datos = [
            proveedor.nit,
            proveedor.empresa,
            proveedor.contacto,
            proveedor.correo,
            proveedor.telefono,
            proveedor.ciudad,
            proveedor.direccion
        ]

        for columna, valor in enumerate(
            datos,
            start=1
        ):

            ws.cell(
                row=fila,
                column=columna
            ).value = valor

        fila += 1

    fecha = datetime.now().strftime(
        "%d-%m-%Y"
    )

    response = HttpResponse(
        content_type=
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    response["Content-Disposition"] = (
        f'attachment; filename="Proveedores_{fecha}.xlsx"'
    )

    wb.save(response)

    return response


# ==========================================
# EXPORTAR PROVEEDORES PDF
# ==========================================

def exportar_proveedores_pdf(request):

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        SimpleDocTemplate,
        Table,
        TableStyle,
        Paragraph,
        Spacer
    )
    from datetime import datetime

    # ==========================================
    # RESPUESTA
    # ==========================================

    response = HttpResponse(
        content_type="application/pdf"
    )

    response["Content-Disposition"] = (
        'attachment; filename="Reporte_Proveedores.pdf"'
    )

    # ==========================================
    # DOCUMENTO
    # ==========================================

    documento = SimpleDocTemplate(
        response,
        pagesize=landscape(letter),
        rightMargin=12 * mm,
        leftMargin=12 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm
    )

    # ==========================================
    # ESTILOS
    # ==========================================

    estilos = getSampleStyleSheet()

    titulo = ParagraphStyle(
        "Titulo",
        parent=estilos["Title"],
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=24,
        alignment=TA_CENTER,
        spaceAfter=8
    )

    subtitulo = ParagraphStyle(
        "Subtitulo",
        parent=estilos["Normal"],
        fontName="Helvetica",
        fontSize=10,
        leading=14,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#64748B"),
        spaceAfter=15
    )

    texto_tabla = ParagraphStyle(
        "TextoTabla",
        parent=estilos["Normal"],
        fontName="Helvetica",
        fontSize=8,
        leading=10,
        alignment=TA_LEFT
    )

    texto_centro = ParagraphStyle(
        "TextoCentro",
        parent=texto_tabla,
        alignment=TA_CENTER
    )

    # ==========================================
    # PROVEEDORES
    # ==========================================

    proveedores = Proveedor.objects.all()

    total_proveedores = proveedores.count()

    # ==========================================
    # FECHA
    # ==========================================

    fecha = datetime.now().strftime(
        "%d/%m/%Y %H:%M"
    )

    # ==========================================
    # CONTENIDO
    # ==========================================

    elementos = []

    elementos.append(
        Paragraph(
            "MINING STAR ERP",
            titulo
        )
    )

    elementos.append(
        Paragraph(
            "REPORTE DE PROVEEDORES",
            subtitulo
        )
    )

    elementos.append(
        Paragraph(
            f"<b>Fecha de generación:</b> {fecha}",
            texto_centro
        )
    )

    elementos.append(
        Spacer(1, 8)
    )

    elementos.append(
        Paragraph(
            f"<b>Total de proveedores:</b> {total_proveedores}",
            texto_centro
        )
    )

    elementos.append(
        Spacer(1, 15)
    )

    # ==========================================
    # ENCABEZADOS
    # ==========================================

    datos = [

        [
            Paragraph("<b>NIT</b>", texto_centro),
            Paragraph("<b>Empresa</b>", texto_centro),
            Paragraph("<b>Contacto</b>", texto_centro),
            Paragraph("<b>Correo</b>", texto_centro),
            Paragraph("<b>Teléfono</b>", texto_centro),
            Paragraph("<b>Ciudad</b>", texto_centro),
            Paragraph("<b>Dirección</b>", texto_centro),
        ]

    ]

    # ==========================================
    # FILAS
    # ==========================================

    for proveedor in proveedores:

        datos.append(

            [

                Paragraph(
                    str(proveedor.nit),
                    texto_centro
                ),

                Paragraph(
                    str(proveedor.empresa),
                    texto_tabla
                ),

                Paragraph(
                    str(proveedor.contacto),
                    texto_tabla
                ),

                Paragraph(
                    str(proveedor.correo),
                    texto_tabla
                ),

                Paragraph(
                    str(proveedor.telefono),
                    texto_centro
                ),

                Paragraph(
                    str(proveedor.ciudad),
                    texto_tabla
                ),

                Paragraph(
                    str(proveedor.direccion),
                    texto_tabla
                ),

            ]

        )

    # ==========================================
    # TABLA
    # ==========================================

    tabla = Table(
        datos,
        repeatRows=1,
        colWidths=[
            30 * mm,
            45 * mm,
            40 * mm,
            55 * mm,
            32 * mm,
            35 * mm,
            55 * mm,
        ]
    )

    # ==========================================
    # ESTILO TABLA
    # ==========================================

    estilo_tabla = [

        (
            "BACKGROUND",
            (0, 0),
            (-1, 0),
            colors.HexColor("#1F4E78")
        ),

        (
            "TEXTCOLOR",
            (0, 0),
            (-1, 0),
            colors.white
        ),

        (
            "FONTNAME",
            (0, 0),
            (-1, 0),
            "Helvetica-Bold"
        ),

        (
            "VALIGN",
            (0, 0),
            (-1, -1),
            "MIDDLE"
        ),

        (
            "GRID",
            (0, 0),
            (-1, -1),
            0.5,
            colors.HexColor("#CBD5E1")
        ),

        (
            "ROWBACKGROUNDS",
            (0, 1),
            (-1, -1),
            [
                colors.white,
                colors.HexColor("#F8FAFC")
            ]
        ),

        (
            "TOPPADDING",
            (0, 0),
            (-1, -1),
            7
        ),

        (
            "BOTTOMPADDING",
            (0, 0),
            (-1, -1),
            7
        ),

        (
            "LEFTPADDING",
            (0, 0),
            (-1, -1),
            5
        ),

        (
            "RIGHTPADDING",
            (0, 0),
            (-1, -1),
            5
        ),

    ]

    tabla.setStyle(
        TableStyle(estilo_tabla)
    )

    elementos.append(tabla)

    elementos.append(
        Spacer(1, 15)
    )

    # ==========================================
    # PIE DEL REPORTE
    # ==========================================

    elementos.append(
        Paragraph(
            "Mining Star ERP | Sistema Integral de Gestión Minera",
            subtitulo
        )
    )

    # ==========================================
    # GENERAR PDF
    # ==========================================

    documento.build(elementos)

    return response

 
# ==========================================
# VENTAS
# ==========================================

def ventas(request):

    from django.db.models import Sum
    from django.http import JsonResponse
    from django.shortcuts import get_object_or_404, redirect
    from .models import Cliente, Producto, Venta, DetalleVenta

    if request.method == "POST":

        cliente_id = request.POST.get("cliente")
        producto_id = request.POST.get("producto")
        cantidad = int(request.POST.get("cantidad", 0) or 0)

        cliente = get_object_or_404(Cliente, id=cliente_id)
        producto = get_object_or_404(Producto, id=producto_id)

        precio = producto.precio_venta
        subtotal = cantidad * precio

        venta = Venta.objects.create(cliente=cliente, total=subtotal)

        DetalleVenta.objects.create(
            venta=venta,
            producto=producto,
            cantidad=cantidad,
            precio_unitario=precio,
            subtotal=subtotal
        )

        producto.stock = max(producto.stock - cantidad, 0)
        producto.save()

        if request.headers.get("X-Requested-With") == "XMLHttpRequest":
            return JsonResponse({
                "success": True,
                "message": "Venta registrada correctamente."
            })

        messages.success(request, "Venta registrada correctamente.")
        return redirect("ventas")

    clientes = Cliente.objects.all()
    productos = Producto.objects.all()

    ventas = (
        Venta.objects
        .select_related("cliente")
        .prefetch_related("detalles__producto")
        .order_by("-fecha")[:20]
    )

    total_vendido = (
        Venta.objects.aggregate(total=Sum("total"))["total"] or 0
    )

    return render(
        request,
        "ventas/ventas.html",
        {
            "clientes": clientes,
            "productos": productos,
            "ventas": ventas,
            "total_vendido": total_vendido,
        }
    )
# ==========================================
# ELIMINAR VENTA
# ==========================================

def eliminar_venta(request, id):

    from django.shortcuts import get_object_or_404, redirect
    from .models import Venta

    venta = get_object_or_404(Venta, id=id)

    # Devolvemos el stock de los productos antes de borrar
    for detalle in venta.detalles.all():
        producto = detalle.producto
        producto.stock += detalle.cantidad
        producto.save()

    venta.delete()

    messages.success(request, "Venta eliminada correctamente.")

    return redirect("ventas")